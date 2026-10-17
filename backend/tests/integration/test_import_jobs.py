"""Catalog import (19 §6, US-M04-04 catalog part, QT-NFR-SIGNUP-01 catalog half) — real app, DB, Redis, SAQ."""

from __future__ import annotations

import io
import time
from typing import Any

import pytest
from openpyxl import Workbook
from sqlalchemy import text

from app.core.db.session import TenantContextError, tenant_transaction
from app.modules.access.domain.permissions import SystemRole
from app.modules.imports.tasks import import_jobs_validate
from tests.support.catalog import Env, Shop, catalog_env_fixture, create_product

pytestmark = [pytest.mark.db, pytest.mark.redis]
__all__ = ["catalog_env_fixture"]

API = "/api/v1/import-jobs"
HEADER = "name;name_ru;barcode;unit_code;mxik_code;vat_rate;price;category;is_marked"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def err(response: Any) -> str:
    code: str = response.json()["error"]["code"]
    return code


def csv_file(*rows: str, header: str = HEADER) -> bytes:
    return ("\r\n".join([header, *rows]) + "\r\n").encode("utf-8")


def xlsx_file(rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(HEADER.split(";"))
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def upload(env: Env, shop: Shop, content: bytes, name: str = "products.csv", ctype: str = "text/csv",
                 kind: str = "products", role: str = "owner") -> Any:  # fmt: skip
    return await env.client.post(
        API, data={"kind": kind}, files={"file": (name, content, ctype)}, headers=shop.h(role)
    )


async def job(env: Env, shop: Shop, job_id: str) -> dict[str, Any]:
    r = await env.client.get(f"{API}/{job_id}", headers=shop.h(key=None))
    assert r.status_code == 200, r.text
    body: dict[str, Any] = r.json()
    return body


async def test_template_and_kinds(env: Env) -> None:
    shop = await env.shop("Shablon", SystemRole.OWNER, SystemRole.CASHIER)
    r = await env.client.get(f"{API}/templates/products", headers=shop.h(key=None))
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/csv")
    assert r.headers["content-disposition"] == 'attachment; filename="products.csv"'
    assert r.content.decode("utf-8-sig").strip() == HEADER
    for kind in ("opening_debts",):  # Bosqich 4 (opening_stock — Bosqich 2, test_stock_openings.py)
        refused = await env.client.get(f"{API}/templates/{kind}", headers=shop.h(key=None))
        assert refused.status_code == 422
        assert refused.json()["error"]["details"][0]["message_key"] == "validation.import_kind_not_supported"
        assert (await upload(env, shop, csv_file(), kind=kind)).status_code == 422
    opening = await env.client.get(f"{API}/templates/opening_stock", headers=shop.h(key=None))
    assert opening.status_code == 200
    assert opening.content.decode("utf-8-sig").strip() == "barcode;qty;unit_cost;expiry_date"
    assert (await upload(env, shop, csv_file(), kind="nonsense")).status_code == 422
    denied = await env.client.get(f"{API}/templates/products", headers=shop.h("cashier", key=None))
    assert denied.status_code == 403


async def test_upload_limits_and_file_checks(env: Env) -> None:
    shop = await env.shop("Fayl tekshiruvi")
    wrong_type = await upload(env, shop, b"%PDF-1.4", name="katalog.pdf", ctype="application/pdf")
    assert wrong_type.status_code == 415 and err(wrong_type) == "UNSUPPORTED_FILE_TYPE"
    no_columns = await upload(env, shop, b"nom;narx\r\nA;1\r\n")
    assert no_columns.status_code == 422 and err(no_columns) == "IMPORT_FILE_INVALID"
    bad_encoding = await upload(env, shop, b"name;unit_code;vat_rate\r\n\xd1\xee\xea;box;12\r\n")
    assert bad_encoding.status_code == 422 and err(bad_encoding) == "IMPORT_FILE_INVALID"
    fake_xlsx = await upload(env, shop, b"not a workbook", name="k.xlsx", ctype=XLSX)
    assert fake_xlsx.status_code == 422 and err(fake_xlsx) == "IMPORT_FILE_INVALID"
    # SEC-05: a few-KB workbook with one cell at XFD1048576 is refused fast, not padded for minutes
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(HEADER.split(";"))
    ws["XFD1048576"] = "x"
    buf = io.BytesIO()
    wb.save(buf)
    started = time.perf_counter()
    far_cell = await upload(env, shop, buf.getvalue(), name="k.xlsx", ctype=XLSX)
    assert far_cell.status_code == 413 and err(far_cell) == "FILE_TOO_LARGE"
    assert far_cell.json()["error"]["params"] == {"max_rows": env.settings.import_max_rows}
    assert time.perf_counter() - started < 5
    # SEC-13: JSON bodies are capped by the body-limit middleware before routing
    huge_json = await env.client.post(f"{API}/00000000-0000-0000-0000-000000000000/commit",
                                      json={"pad": "x" * (2 * 1024 * 1024)}, headers=shop.h())  # fmt: skip
    assert huge_json.status_code == 413 and err(huge_json) == "FILE_TOO_LARGE"
    env.app.state.imports._limits = env.app.state.imports._limits.__class__(max_bytes=2048, max_rows=3)
    too_big = await upload(env, shop, csv_file(*(["A;;;box;;12;;;"] * 400)))
    assert too_big.status_code == 413 and err(too_big) == "FILE_TOO_LARGE"
    too_many = await upload(env, shop, csv_file(*(["A;;;box;;12;;;"] * 4)))
    assert too_many.status_code == 413 and err(too_many) == "FILE_TOO_LARGE"


async def test_us_m04_04_csv_upload_validate_commit(env: Env) -> None:
    """Excel/CSV path of US-M04-04: row/field errors in import_jobs.errors, then commit of the valid rows."""
    shop = await env.shop("Import CSV", SystemRole.OWNER)
    existing = await create_product(env.client, shop, "Bor tovar", barcode="4780000000999")
    content = csv_file(
        "Coca-Cola 1.5L;Кока-кола 1.5Л;4780000000012;box;02202001001000000;12;12000;Ichimliklar;0",
        "Pepsi 1L;;4780000000029;BOX;;12.00;9500.50;ichimliklar;",
        "Non;;;kg;;0;;Non mahsulotlari;yo'q",
        "Suv 0.5;;4780000000036;box;;12;3000;;ha",
        "Takror;;4780000000012;box;;12;1;;",  # duplicate inside the file
        "Bor;;4780000000999;box;;12;1;;",  # already in the catalog
        ";;;tonna;123;15;-5;;?",  # many field errors
    )
    created = await upload(env, shop, content)
    assert created.status_code == 201, created.text
    assert created.json()["status"] == "uploaded" and created.json()["total_rows"] == 0
    job_id = created.json()["id"]

    early = await env.client.post(f"{API}/{job_id}/commit", json={}, headers=shop.h())
    assert early.status_code == 409 and err(early) == "IMPORT_JOB_NOT_VALIDATED"

    v = await env.client.post(f"{API}/{job_id}/validate", headers=shop.h())
    assert v.status_code == 202 and v.json() == {"job_id": job_id, "status": "validating"}
    assert (await job(env, shop, job_id))["status"] == "validating"
    assert await env.drain() == 1
    state = await job(env, shop, job_id)
    assert state["status"] == "validated"
    assert (state["total_rows"], state["ok_rows"], state["error_rows"]) == (7, 4, 3)
    assert state["result"] == {}

    errors = (
        await env.client.get(f"{API}/{job_id}/errors", params={"limit": 3}, headers=shop.h(key=None))
    ).json()
    assert len(errors["items"]) == 3 and errors["next_cursor"]
    rest = await env.client.get(
        f"{API}/{job_id}/errors",
        params={"cursor": errors["next_cursor"], "limit": 50},
        headers=shop.h(key=None),
    )
    all_errors = errors["items"] + rest.json()["items"]
    assert rest.json()["next_cursor"] is None
    by_row = {(e["row"], e["field"], e["code"]) for e in all_errors}
    assert (6, "barcode", "BARCODE_DUPLICATE") in by_row and (7, "barcode", "BARCODE_DUPLICATE") in by_row
    assert (8, "mxik_code", "MXIK_CODE_INVALID") in by_row and (8, "name", "VALIDATION_ERROR") in by_row
    assert all(e["message"].startswith("validation.") for e in all_errors)

    strict = await env.client.post(
        f"{API}/{job_id}/commit", json={"skip_error_rows": False}, headers=shop.h()
    )
    assert strict.status_code == 422
    c = await env.client.post(f"{API}/{job_id}/commit", json={"skip_error_rows": True}, headers=shop.h())
    assert c.status_code == 202 and c.json() == {"job_id": job_id, "status": "committing"}
    assert await env.drain() == 1
    done = await job(env, shop, job_id)
    assert done["status"] == "committed" and done["result"]["created_product_count"] == 4
    assert len(done["result"]["created_doc_ids"]) == 1  # one posted price_changes doc with the first prices
    again = await env.client.post(f"{API}/{job_id}/commit", json={}, headers=shop.h())
    assert again.status_code == 409 and err(again) == "IMPORT_JOB_ALREADY_COMMITTED"
    revalidate = await env.client.post(f"{API}/{job_id}/validate", headers=shop.h())
    assert revalidate.status_code == 409 and err(revalidate) == "IMPORT_JOB_ALREADY_COMMITTED"

    products = (
        await env.client.get("/api/v1/products", params={"size": 50}, headers=shop.h(key=None))
    ).json()
    by_name = {p["name"]: p for p in products["items"]}
    assert set(by_name) == {"Bor tovar", "Coca-Cola 1.5L", "Pepsi 1L", "Non", "Suv 0.5"}
    cola = by_name["Coca-Cola 1.5L"]
    assert cola["name_ru"] == "Кока-кола 1.5Л" and cola["mxik_code"] == "02202001001000000"
    assert [b["barcode"] for b in cola["barcodes"]] == ["4780000000012"]
    assert by_name["Pepsi 1L"]["category_id"] == cola["category_id"] is not None  # "ichimliklar" matched
    assert by_name["Non"]["vat_rate"] == "0.00" and by_name["Non"]["unit_code"] == "kg"
    assert by_name["Suv 0.5"]["is_marked"] is True and by_name["Suv 0.5"]["marking_group_code"] == "general"
    prices = await env.client.get(f"/api/v1/products/{cola['id']}/prices", headers=shop.h(key=None))
    assert prices.json()["items"][0]["price"] == "12000.0000"
    assert prices.json()["items"][0]["price_change_doc_id"] == done["result"]["created_doc_ids"][0]
    assert existing["barcodes"][0]["barcode"] == "4780000000999"
    async with env.db.session_factory() as s, tenant_transaction(s, shop.tenant_id):
        doc = (
            await s.execute(
                text("SELECT status, number FROM app.price_changes WHERE id = :id"),
                {"id": done["result"]["created_doc_ids"][0]},
            )
        ).one()
        labels = (
            await s.execute(text("SELECT count(*) FROM app.label_print_queue WHERE reason = 'new'"))
        ).scalar_one()
        attachment = (
            await s.execute(text("SELECT kind, content_type, size_bytes FROM app.attachments"))
        ).one()
    assert doc[0] == "posted" and doc[1].startswith("C-01-")
    assert labels == 5  # 4 imported + 1 created by hand
    assert attachment == ("import", "text/csv", len(content))


async def test_commit_rechecks_catalog_and_rls(env: Env) -> None:
    a = await env.shop("Import A")
    b = await env.shop("Import B")
    created = await upload(env, a, csv_file("Olma;;4780000000500;kg;;0;15000;Meva;0"))
    job_id = created.json()["id"]
    await env.client.post(f"{API}/{job_id}/validate", headers=a.h())
    await env.drain()
    # the barcode got taken between validate and commit → the row is skipped at commit, not duplicated
    await create_product(env.client, a, "Olma (qo'lda)", barcode="4780000000500")
    for method, url in (("GET", f"{API}/{job_id}"), ("GET", f"{API}/{job_id}/errors"),
                        ("POST", f"{API}/{job_id}/validate"), ("POST", f"{API}/{job_id}/commit")):  # fmt: skip
        r = await env.client.request(method, url, headers=b.h())
        assert r.status_code == 404 and r.json()["error"]["params"]["resource"] == "import_job", url
    await env.client.post(f"{API}/{job_id}/commit", json={}, headers=a.h())
    await env.drain()
    state = await job(env, a, job_id)
    assert state["status"] == "committed" and state["result"]["created_product_count"] == 0
    assert state["result"]["created_doc_ids"] == []
    errors = (await env.client.get(f"{API}/{job_id}/errors", headers=a.h(key=None))).json()["items"]
    assert errors == [
        {"row": 2, "field": "barcode", "code": "BARCODE_DUPLICATE", "message": "validation.barcode_duplicate"}
    ]


async def test_all_rows_invalid_fails_and_job_requires_tenant(env: Env) -> None:
    shop = await env.shop("Hammasi xato")
    created = await upload(env, shop, csv_file(";;;tonna;;12;;;"))
    job_id = created.json()["id"]
    await env.client.post(f"{API}/{job_id}/validate", headers=shop.h())
    await env.drain()
    state = await job(env, shop, job_id)
    assert state["status"] == "failed" and state["ok_rows"] == 0 and state["error_rows"] == 1
    refused = await env.client.post(f"{API}/{job_id}/commit", json={}, headers=shop.h())
    assert refused.status_code == 409 and err(refused) == "IMPORT_JOB_NOT_VALIDATED"
    # a job without tenant fails before touching the database (23 §4)
    with pytest.raises(TenantContextError):
        await import_jobs_validate({"db": env.db}, tenant_id=None, job_id=job_id)
    # a job for another tenant cannot see this row (RLS) and does nothing
    other = await env.shop("Boshqa")
    skipped = await import_jobs_validate({"db": env.db, "settings": env.settings}, tenant_id=str(other.tenant_id),
                                         job_id=job_id)  # fmt: skip
    assert skipped == {"skipped": True}


async def test_qt_nfr_signup_01_catalog_import_500_products_xlsx(env: Env) -> None:
    """QT-NFR-SIGNUP-01 (catalog half): 500-product Excel with prices, upload → validate → commit."""
    shop = await env.shop("Yangi do'kon")
    rows = [
        [f"Tovar {i:03d}", f"Товар {i:03d}", f"478{i:010d}", "box" if i % 5 else "kg", "", 12 if i % 3 else 0,
         1000 + i * 10.5, f"Kategoriya {i % 10}", 0]
        for i in range(500)
    ]  # fmt: skip
    content = xlsx_file(rows)
    started = time.perf_counter()
    created = await upload(env, shop, content, name="katalog.xlsx", ctype=XLSX)
    assert created.status_code == 201, created.text
    job_id = created.json()["id"]
    assert (await env.client.post(f"{API}/{job_id}/validate", headers=shop.h())).status_code == 202
    await env.drain()
    validated = await job(env, shop, job_id)
    assert validated["status"] == "validated" and validated["ok_rows"] == 500, validated
    assert (await env.client.post(f"{API}/{job_id}/commit", json={}, headers=shop.h())).status_code == 202
    await env.drain()
    elapsed = time.perf_counter() - started
    done = await job(env, shop, job_id)
    assert done["status"] == "committed" and done["result"]["created_product_count"] == 500
    print(f"\nQT-NFR-SIGNUP-01 catalog import: 500 products in {elapsed:.2f}s")
    assert elapsed < 120  # budget: signup→first sale ≤ 10 min overall (observed, not blocking)
    async with env.db.session_factory() as s, tenant_transaction(s, shop.tenant_id):
        counts = (
            await s.execute(
                text(
                    "SELECT (SELECT count(*) FROM app.products), (SELECT count(*) FROM app.product_barcodes), "
                    "(SELECT count(*) FROM app.product_prices), (SELECT count(*) FROM app.product_categories)"
                )
            )
        ).one()
    assert tuple(counts) == (500, 500, 500, 10)


async def test_opening_stock_import_needs_the_default_branch_in_scope(env: Env) -> None:
    """SEC-02: the ``opening_stock`` import posts into the default branch's warehouse — a user restricted
    to another branch is refused (upload / validate / commit → 403), and so is the commit job itself."""
    from app.modules.accounting.wiring import business_today
    from app.modules.imports.tasks import import_jobs_commit
    from tests.support.inventory import add_branch, add_user, set_go_live

    shop = await env.shop("Import scope")
    await set_go_live(env, shop, business_today().replace(day=1))
    await create_product(env.client, shop, "Pepsi 1L", barcode="4870000000024")
    branch_b, _ = await add_branch(env, shop, "B2")
    await add_user(env, shop, "accountant_b", SystemRole.ACCOUNTANT, (branch_b,))
    content = b"barcode;qty;unit_cost;expiry_date\r\n4870000000024;5;1000;\r\n"
    refused = await upload(env, shop, content, "opening.csv", kind="opening_stock", role="accountant_b")
    assert refused.status_code == 403 and err(refused) == "BRANCH_SCOPE_DENIED"
    job_id = (await upload(env, shop, content, "opening.csv", kind="opening_stock")).json()["id"]
    denied = await env.client.post(f"{API}/{job_id}/validate", headers=shop.h("accountant_b"))
    assert denied.status_code == 403 and err(denied) == "BRANCH_SCOPE_DENIED"
    await env.client.post(f"{API}/{job_id}/validate", headers=shop.h())
    await env.drain()
    assert (await job(env, shop, job_id))["status"] == "validated"
    denied = await env.client.post(f"{API}/{job_id}/commit", json={}, headers=shop.h("accountant_b"))
    assert denied.status_code == 403 and err(denied) == "BRANCH_SCOPE_DENIED"
    await env.drain()
    assert (await job(env, shop, job_id))["status"] == "validated"
    # the worker re-checks the committing actor's scope (a job enqueued with a restricted scope)
    async with env.db.session_factory() as s, tenant_transaction(s, shop.tenant_id):
        await s.execute(
            text("UPDATE app.import_jobs SET status = 'committing' WHERE id = :id"), {"id": job_id}
        )
    result = await import_jobs_commit(
        {"db": env.db, "settings": env.settings}, tenant_id=str(shop.tenant_id), job_id=job_id,
        user_id=str(shop.user_ids["accountant_b"]), branch_ids=[str(branch_b)],
    )  # fmt: skip
    assert result == {"status": "failed", "code": "BRANCH_SCOPE_DENIED"}
    state = await job(env, shop, job_id)
    assert state["status"] == "failed" and state["result"] == {}
    async with env.db.session_factory() as s, tenant_transaction(s, shop.tenant_id):
        count = (await s.execute(text("SELECT count(*) FROM app.stock_openings"))).scalar()
    assert count == 0
