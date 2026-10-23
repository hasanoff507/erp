"""products.csv/xlsx parsing and row validation (US-M04-04, 19 §6)."""

from __future__ import annotations

import io
import time
import zipfile
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import Workbook

from app.modules.imports.application.parsing import (
    MAX_COLUMNS,
    MAX_COMPRESSION_RATIO,
    MAX_UNPACKED_BYTES,
    FileInvalidError,
    FileTooLargeError,
    TooManyRowsError,
    check_xlsx_zip,
    read_table,
)
from app.modules.imports.domain.products import REQUIRED_COLUMNS, cell_text, parse_product_row, template_csv

UNITS = {"pcs", "kg", "l", "m", "box"}


def test_template_has_the_documented_columns() -> None:
    assert (
        template_csv("products").strip()
        == "name;name_ru;barcode;unit_code;mxik_code;vat_rate;price;category;is_marked"
    )
    with pytest.raises(ValueError, match="opening_stock"):
        template_csv("opening_stock")


def test_csv_semicolon_utf8_bom() -> None:
    content = "﻿name;unit_code;vat_rate;price\r\nСок;pcs;12;9000.50\r\n;;;\r\nNon;kg;0;\r\n".encode()
    table = read_table(content, "csv", required=REQUIRED_COLUMNS, max_rows=10)
    assert [r[0] for r in table.rows] == [2, 4]
    assert table.rows[0][1]["name"] == "Сок"


@pytest.mark.parametrize(
    ("content", "reason"),
    [
        (b"name;unit_code\nA;pcs\n", "missing_columns"),
        (b"name;unit_code;vat_rate\n\xd1\xee\xea;pcs;12\n", "encoding"),  # cp1251 "Сок"
        (b"", "empty"),
    ],
)
def test_csv_invalid(content: bytes, reason: str) -> None:
    with pytest.raises(FileInvalidError, match=reason):
        read_table(content, "csv", required=REQUIRED_COLUMNS, max_rows=10)


def test_row_limit() -> None:
    content = ("name;unit_code;vat_rate\n" + "A;pcs;12\n" * 4).encode()
    with pytest.raises(TooManyRowsError):
        read_table(content, "csv", required=REQUIRED_COLUMNS, max_rows=3)


def test_xlsx_numeric_cells_and_bad_zip() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Name", "Unit_Code", "VAT_rate", "MXIK_code", "barcode", "price"])
    ws.append(["Pepsi", "pcs", 12, 2202001002000000, 4780000000029, 7500.5])
    buf = io.BytesIO()
    wb.save(buf)
    table = read_table(buf.getvalue(), "xlsx", required=REQUIRED_COLUMNS, max_rows=10)
    row = table.rows[0][1]
    assert row["mxik_code"] == "02202001002000000"  # Excel dropped the leading zero
    assert row["barcode"] == "4780000000029" and row["price"] == "7500.5" and row["vat_rate"] == "12"
    with pytest.raises(FileInvalidError):
        read_table(b"PK\x03\x04garbage", "xlsx", required=REQUIRED_COLUMNS, max_rows=10)
    with pytest.raises(FileInvalidError):
        read_table(b"not a zip", "xlsx", required=REQUIRED_COLUMNS, max_rows=10)


def test_row_validation() -> None:
    ok, errors = parse_product_row(
        2,
        {"name": "Coca-Cola", "unit_code": "PCS", "vat_rate": "12%", "price": "12000", "is_marked": "ha",
         "mxik_code": "02202001001000000", "category": "Ichimliklar", "name_ru": "Кока-кола"},
        UNITS,
    )  # fmt: skip
    assert errors == [] and ok is not None
    assert ok.unit_code == "pcs" and ok.vat_rate == Decimal("12.00") and ok.price == Decimal("12000.0000")
    assert ok.is_marked and ok.name_ru == "Кока-кола"

    bad, errors = parse_product_row(
        3,
        {
            "name": "",
            "unit_code": "tonna",
            "vat_rate": "15",
            "price": "-1",
            "mxik_code": "123",
            "is_marked": "?",
        },
        UNITS,
    )
    assert bad is None
    fields = {(e.field, e.code) for e in errors}
    assert ("name", "VALIDATION_ERROR") in fields and ("unit_code", "VALIDATION_ERROR") in fields
    assert ("vat_rate", "VALIDATION_ERROR") in fields and ("mxik_code", "MXIK_CODE_INVALID") in fields
    assert ("price", "VALIDATION_ERROR") in fields and ("is_marked", "VALIDATION_ERROR") in fields
    assert all(e.row == 3 and e.message.startswith("validation.") for e in errors)
    assert cell_text(12.0, "price") == "12" and cell_text(None, "x") == "" and cell_text(True, "x") == "1"


def _xlsx(build: Any) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["name", "unit_code", "vat_rate"])
    build(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_far_cell_fails_fast() -> None:
    """SEC-05: one cell at XFD1048576 (a few KB) must not pad a million rows x 16 384 columns."""

    def far(ws: Any) -> None:
        ws.append(["Pepsi", "pcs", 12])
        ws["XFD1048576"] = "x"

    content = _xlsx(far)
    assert len(content) < 20_000
    started = time.perf_counter()
    with pytest.raises(TooManyRowsError) as info:
        read_table(content, "xlsx", required=REQUIRED_COLUMNS, max_rows=20_000)
    assert time.perf_counter() - started < 2.0
    assert info.value.params == {"max_rows": 20_000}


def test_blank_lines_count_against_the_hard_cap() -> None:
    content = ("name;unit_code;vat_rate\n" + ";;\n" * 200 + "A;pcs;12\n").encode()
    with pytest.raises(TooManyRowsError):
        read_table(content, "csv", required=REQUIRED_COLUMNS, max_rows=10)  # cap = 10*2+100 lines


def test_wide_header_reads_only_the_first_columns() -> None:
    def wide(ws: Any) -> None:
        ws.cell(row=1, column=200, value="far_column")
        ws.append(["Pepsi", "pcs", 12])

    table = read_table(_xlsx(wide), "xlsx", required=REQUIRED_COLUMNS, max_rows=10)
    assert len(table.header) == MAX_COLUMNS
    assert table.rows[0][1] == {"name": "Pepsi", "unit_code": "pcs", "vat_rate": "12"}


def _zip(name: str, data: bytes) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(name, data)
    return buf.getvalue()


def test_zip_bomb_ratio_and_unpacked_size_are_refused_before_parsing() -> None:
    bomb = _zip("xl/worksheets/sheet1.xml", b"\0" * (8 * 1024 * 1024))  # ~1000:1
    started = time.perf_counter()
    with pytest.raises(FileTooLargeError) as info:
        read_table(bomb, "xlsx", required=REQUIRED_COLUMNS, max_rows=10)
    assert info.value.params == {"max_compression_ratio": MAX_COMPRESSION_RATIO}
    assert time.perf_counter() - started < 2.0
    huge = _zip("xl/sharedStrings.xml", bytes(MAX_UNPACKED_BYTES + 1))  # size wins over ratio
    with pytest.raises(FileTooLargeError) as info:
        check_xlsx_zip(huge)
    assert info.value.params == {"max_unpacked_bytes": MAX_UNPACKED_BYTES}
