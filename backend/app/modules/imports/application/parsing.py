"""CSV/XLSX → header + rows (17 §12, 19 §6). CPU-bound; callers run it in a worker thread.

CSV: UTF-8 (BOM allowed), ``;`` separator (``,`` accepted when the header has no ``;``), decimal point.
XLSX: first worksheet, first row = header. Fully empty rows are skipped; ``row`` is the sheet line number.

Hostile files (SEC-05): the XLSX zip is inspected before openpyxl touches it (total unpacked size
and compression ratio), only the first ``MAX_COLUMNS`` columns are read, and *every* line — blank
ones included — counts against a hard cap, so a tiny file with one cell at ``XFD1048576`` or a
zip bomb fails fast with ``413 FILE_TOO_LARGE`` instead of padding a million rows.
"""

from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass
from typing import Any, Final, Literal

from openpyxl import load_workbook

from app.modules.imports.domain.products import cell_text

FileFormat = Literal["csv", "xlsx"]
XLSX_MAGIC: Final = b"PK\x03\x04"
MAX_COLUMNS: Final = 64  # the products template has 9 columns; the rest is slack for user extras
MAX_UNPACKED_BYTES: Final = 50 * 1024 * 1024
MAX_COMPRESSION_RATIO: Final = 100
RATIO_CHECK_MIN_BYTES: Final = 1024 * 1024  # tiny files may compress oddly well; they are harmless


class FileInvalidError(Exception):
    """Format/encoding error → ``422 IMPORT_FILE_INVALID``."""


class FileTooLargeError(Exception):
    """Size limit exceeded (rows, unpacked bytes, compression ratio) → ``413 FILE_TOO_LARGE``."""

    def __init__(self, params: dict[str, int]) -> None:
        super().__init__(", ".join(f"{k}={v}" for k, v in params.items()))
        self.params = params


class TooManyRowsError(FileTooLargeError):
    """Row limit exceeded → ``413 FILE_TOO_LARGE`` with ``params.max_rows``."""

    def __init__(self, max_rows: int) -> None:
        super().__init__({"max_rows": max_rows})


def hard_line_cap(max_rows: int) -> int:
    """Every line read (blank ones too) counts against this; data rows against ``max_rows``."""
    return max_rows * 2 + 100


@dataclass(frozen=True, slots=True)
class Table:
    header: tuple[str, ...]
    rows: list[tuple[int, dict[str, str]]]


def _normalize_header(cells: list[Any]) -> tuple[str, ...]:
    return tuple(str(c).strip().lower().lstrip("﻿") if c is not None else "" for c in cells)


def _csv_rows(content: bytes) -> Iterator[tuple[int, list[Any]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise FileInvalidError("encoding") from exc
    first = text.split("\n", 1)[0]
    delimiter = ";" if ";" in first or "," not in first else ","
    try:
        for index, cells in enumerate(
            csv.reader(io.StringIO(text, newline=""), delimiter=delimiter), start=1
        ):
            yield index, list(cells)
    except csv.Error as exc:
        raise FileInvalidError("csv") from exc


def check_xlsx_zip(content: bytes) -> None:
    """Zip-bomb guard from the central directory alone (nothing is decompressed)."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
    except (zipfile.BadZipFile, ValueError, OSError) as exc:
        raise FileInvalidError("xlsx") from exc
    unpacked = sum(info.file_size for info in infos)
    packed = sum(info.compress_size for info in infos)
    if unpacked > MAX_UNPACKED_BYTES:
        raise FileTooLargeError({"max_unpacked_bytes": MAX_UNPACKED_BYTES})
    if unpacked > RATIO_CHECK_MIN_BYTES and unpacked > packed * MAX_COMPRESSION_RATIO:
        raise FileTooLargeError({"max_compression_ratio": MAX_COMPRESSION_RATIO})


def _xlsx_rows(content: bytes, max_lines: int) -> Iterator[tuple[int, list[Any]]]:
    if not content.startswith(XLSX_MAGIC):
        raise FileInvalidError("xlsx_magic")
    check_xlsx_zip(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise FileInvalidError("xlsx") from exc
    try:
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            raise FileInvalidError("xlsx_empty")
        # max_col: no padding to a far <dimension>; max_row: the far tail is never generated
        rows = sheet.iter_rows(max_row=max_lines + 1, max_col=MAX_COLUMNS, values_only=True)
        for index, cells in enumerate(rows, start=1):
            yield index, list(cells)
    finally:
        workbook.close()


def read_table(content: bytes, fmt: FileFormat, *, required: frozenset[str], max_rows: int) -> Table:
    max_lines = hard_line_cap(max_rows)
    rows_iter = _csv_rows(content) if fmt == "csv" else _xlsx_rows(content, max_lines)
    header: tuple[str, ...] | None = None
    columns: list[tuple[int, str]] = []
    rows: list[tuple[int, dict[str, str]]] = []
    for line_no, cells in rows_iter:
        if line_no > max_lines:
            raise TooManyRowsError(max_rows)
        if header is None:
            if not any(cell_text(c, "") for c in cells):
                continue
            header = _normalize_header(cells)
            missing = required - set(header)
            if missing:
                raise FileInvalidError("missing_columns:" + ",".join(sorted(missing)))
            columns = [(i, column) for i, column in enumerate(header) if column]
            continue
        values = {column: cell_text(cells[i] if i < len(cells) else None, column) for i, column in columns}
        if not any(values.values()):
            continue
        rows.append((line_no, values))
        if len(rows) > max_rows:
            raise TooManyRowsError(max_rows)
    if header is None:
        raise FileInvalidError("empty")
    return Table(header, rows)
